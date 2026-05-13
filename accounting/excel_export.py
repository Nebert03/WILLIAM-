"""Export financial reports to a styled Excel workbook."""
from __future__ import annotations
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── colour palette (matches Invegrow brand) ──────────────────────────────────
_GREEN_DARK  = "1B5E20"
_GREEN_MID   = "388E3C"
_GREEN_LIGHT = "C8E6C9"
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_TITLE_FONT  = Font(name="Calibri", bold=True, size=14, color=_GREEN_DARK)
_LABEL_FONT  = Font(name="Calibri", bold=True, size=10)
_BODY_FONT   = Font(name="Calibri", size=10)
_SECTION_FILL = PatternFill("solid", fgColor=_GREEN_LIGHT)
_HEADER_FILL  = PatternFill("solid", fgColor=_GREEN_DARK)
_SUBHDR_FILL  = PatternFill("solid", fgColor=_GREEN_MID)
_THIN = Side(style="thin")
_BORDER = Border(bottom=_THIN)
_DOUBLE_BORDER = Border(top=_THIN, bottom=Side(style="double"))
_MONEY = '#,##0.00'
_CENTER = Alignment(horizontal="center")
_RIGHT  = Alignment(horizontal="right")


def _col(ws, col: int, width: float):
    ws.column_dimensions[get_column_letter(col)].width = width


def _hrow(ws, row: int, *values, fill=None, font=None, number_fmt=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        if fill:
            cell.fill = fill
        if font:
            cell.font = font
        if number_fmt and isinstance(v, (int, float)):
            cell.number_format = number_fmt
    return row + 1


def export_workbook(
    trial_balance: dict,
    income_statement: dict,
    balance_sheet: dict,
    output_path: str | Path,
) -> Path:
    output_path = Path(output_path)
    wb = openpyxl.Workbook()

    _write_trial_balance(wb.active, trial_balance)
    wb.active.title = "Trial Balance"

    ws_is = wb.create_sheet("Income Statement")
    _write_income_statement(ws_is, income_statement)

    ws_bs = wb.create_sheet("Balance Sheet")
    _write_balance_sheet(ws_bs, balance_sheet)

    ws_coa = wb.create_sheet("Chart of Accounts")
    _write_coa(ws_coa, trial_balance["rows"])

    wb.save(output_path)
    return output_path


# ── Trial Balance sheet ───────────────────────────────────────────────────────

def _write_trial_balance(ws, tb: dict):
    _col(ws, 1, 10); _col(ws, 2, 32); _col(ws, 3, 16); _col(ws, 4, 16)

    r = 1
    ws.cell(r, 1, tb["title"]).font = _TITLE_FONT
    ws.merge_cells(f"A{r}:D{r}")
    r += 1
    ws.cell(r, 1, tb["period"]).font = _BODY_FONT
    ws.merge_cells(f"A{r}:D{r}")
    r += 2

    r = _hrow(ws, r, "Code", "Account Name", "Debit", "Credit",
              fill=_HEADER_FILL, font=_HEADER_FONT)

    prev_type = None
    for row in tb["rows"]:
        if row["type"] != prev_type:
            ws.cell(r, 1, row["type"]).fill  = _SECTION_FILL
            ws.cell(r, 1).font = _LABEL_FONT
            ws.merge_cells(f"A{r}:D{r}")
            r += 1
            prev_type = row["type"]

        ws.cell(r, 1, row["code"]).font = _BODY_FONT
        ws.cell(r, 2, row["name"]).font = _BODY_FONT
        dr_cell = ws.cell(r, 3, row["debit"] if row["debit"] else None)
        cr_cell = ws.cell(r, 4, row["credit"] if row["credit"] else None)
        dr_cell.number_format = _MONEY
        cr_cell.number_format = _MONEY
        dr_cell.alignment = _RIGHT
        cr_cell.alignment = _RIGHT
        r += 1

    # totals row
    for c in range(1, 5):
        ws.cell(r, c).border = _DOUBLE_BORDER
    ws.cell(r, 2, "TOTALS").font = _LABEL_FONT
    t_dr = ws.cell(r, 3, tb["debit_total"])
    t_cr = ws.cell(r, 4, tb["credit_total"])
    t_dr.number_format = t_cr.number_format = _MONEY
    t_dr.font = t_cr.font = _LABEL_FONT
    t_dr.alignment = t_cr.alignment = _RIGHT
    r += 1
    status = "BALANCED" if tb["balanced"] else "DOES NOT BALANCE"
    c = ws.cell(r, 2, status)
    c.font = Font(bold=True, color="1B5E20" if tb["balanced"] else "C62828")


# ── Income Statement sheet ────────────────────────────────────────────────────

def _write_income_statement(ws, is_: dict):
    _col(ws, 1, 10); _col(ws, 2, 32); _col(ws, 3, 18)

    r = 1
    ws.cell(r, 1, is_["title"]).font = _TITLE_FONT
    ws.merge_cells(f"A{r}:C{r}")
    r += 1
    ws.cell(r, 1, is_["period"]).font = _BODY_FONT
    ws.merge_cells(f"A{r}:C{r}")
    r += 2

    def _section(label, rows, total_label, total):
        nonlocal r
        ws.cell(r, 1, label).fill = _HEADER_FILL
        ws.cell(r, 1).font = _HEADER_FONT
        ws.merge_cells(f"A{r}:C{r}")
        r += 1
        for row in rows:
            ws.cell(r, 1, row["code"]).font = _BODY_FONT
            ws.cell(r, 2, row["name"]).font = _BODY_FONT
            v = ws.cell(r, 3, row["balance"])
            v.number_format = _MONEY; v.alignment = _RIGHT
            r += 1
        for c in range(1, 4):
            ws.cell(r, c).border = _BORDER
        ws.cell(r, 2, total_label).font = _LABEL_FONT
        tot = ws.cell(r, 3, total)
        tot.number_format = _MONEY; tot.alignment = _RIGHT; tot.font = _LABEL_FONT
        r += 2

    _section("REVENUE",  is_["revenue_rows"],  "Total Revenue",  is_["total_revenue"])
    _section("EXPENSES", is_["expense_rows"],  "Total Expenses", is_["total_expenses"])

    for c in range(1, 4):
        ws.cell(r, c).border = _DOUBLE_BORDER
    label = "NET INCOME" if is_["net_income"] >= 0 else "NET LOSS"
    ws.cell(r, 2, label).font = Font(bold=True, size=11)
    ni = ws.cell(r, 3, is_["net_income"])
    ni.number_format = _MONEY; ni.alignment = _RIGHT
    ni.font = Font(bold=True, size=11,
                   color=_GREEN_DARK if is_["net_income"] >= 0 else "C62828")


# ── Balance Sheet ─────────────────────────────────────────────────────────────

def _write_balance_sheet(ws, bs: dict):
    _col(ws, 1, 10); _col(ws, 2, 34); _col(ws, 3, 18)

    r = 1
    ws.cell(r, 1, bs["title"]).font = _TITLE_FONT
    ws.merge_cells(f"A{r}:C{r}")
    r += 1
    ws.cell(r, 1, f"As of {bs['as_of']}").font = _BODY_FONT
    ws.merge_cells(f"A{r}:C{r}")
    r += 2

    def _section(label, rows, total_label, total, extra_rows=None):
        nonlocal r
        ws.cell(r, 1, label).fill = _HEADER_FILL
        ws.cell(r, 1).font = _HEADER_FONT
        ws.merge_cells(f"A{r}:C{r}")
        r += 1
        for row in rows:
            ws.cell(r, 1, row["code"]).font = _BODY_FONT
            ws.cell(r, 2, row["name"]).font = _BODY_FONT
            v = ws.cell(r, 3, row["balance"])
            v.number_format = _MONEY; v.alignment = _RIGHT
            r += 1
        if extra_rows:
            for (name, val) in extra_rows:
                ws.cell(r, 2, name).font = _BODY_FONT
                v = ws.cell(r, 3, val)
                v.number_format = _MONEY; v.alignment = _RIGHT
                r += 1
        for c in range(1, 4):
            ws.cell(r, c).border = _BORDER
        ws.cell(r, 2, total_label).font = _LABEL_FONT
        tot = ws.cell(r, 3, total)
        tot.number_format = _MONEY; tot.alignment = _RIGHT; tot.font = _LABEL_FONT
        r += 2

    _section("ASSETS",      bs["asset_rows"],     "Total Assets",      bs["total_assets"])
    _section("LIABILITIES", bs["liability_rows"],  "Total Liabilities", bs["total_liabilities"])
    _section("EQUITY",      bs["equity_rows"],    "Total Equity",      bs["total_equity"],
             extra_rows=[("Net Income (Current Period)", bs["net_income"])])

    for c in range(1, 4):
        ws.cell(r, c).border = _DOUBLE_BORDER
    ws.cell(r, 2, "Total Liabilities + Equity").font = Font(bold=True, size=11)
    le = ws.cell(r, 3, bs["total_l_e"])
    le.number_format = _MONEY; le.alignment = _RIGHT; le.font = Font(bold=True, size=11)
    r += 1
    status = "BALANCED" if bs["balanced"] else "DOES NOT BALANCE"
    c = ws.cell(r, 2, status)
    c.font = Font(bold=True, color=_GREEN_DARK if bs["balanced"] else "C62828")


# ── Chart of Accounts sheet ──────────────────────────────────────────────────

def _write_coa(ws, rows: list[dict]):
    _col(ws, 1, 10); _col(ws, 2, 32); _col(ws, 3, 14); _col(ws, 4, 10)
    _col(ws, 5, 14); _col(ws, 6, 14)

    r = 1
    ws.cell(r, 1, "Chart of Accounts").font = _TITLE_FONT
    ws.merge_cells(f"A{r}:F{r}")
    r += 2

    r = _hrow(ws, r, "Code", "Account Name", "Type", "Normal Side", "Debit Total", "Credit Total",
              fill=_HEADER_FILL, font=_HEADER_FONT)
    for row in rows:
        ws.cell(r, 1, row["code"]).font = _BODY_FONT
        ws.cell(r, 2, row["name"]).font = _BODY_FONT
        ws.cell(r, 3, row["type"]).font = _BODY_FONT
        ws.cell(r, 4, row["normal_side"]).font = _BODY_FONT
        dr = ws.cell(r, 5, row["debit"])
        cr = ws.cell(r, 6, row["credit"])
        dr.number_format = cr.number_format = _MONEY
        dr.alignment = cr.alignment = _RIGHT
        r += 1
